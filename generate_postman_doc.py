from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

# ── Styles ──────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(4)

for level in range(1, 4):
    h = doc.styles[f'Heading {level}']
    h.font.name = 'Calibri'
    h.font.color.rgb = RGBColor(0x1A, 0x56, 0xDB)


def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # Header row
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(9)
    # Data rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.add_paragraph()


def add_json_block(label, json_text):
    p = doc.add_paragraph()
    run = p.add_run(f"{label}:")
    run.bold = True
    run.font.size = Pt(10)
    p2 = doc.add_paragraph()
    run2 = p2.add_run(json_text)
    run2.font.name = 'Courier New'
    run2.font.size = Pt(9)
    p2.paragraph_format.left_indent = Cm(1)


BASE = "http://localhost:8080"

# ════════════════════════════════════════════════════════════════════
# TITLE PAGE
# ════════════════════════════════════════════════════════════════════
doc.add_paragraph()
doc.add_paragraph()
title = doc.add_heading('OKR Tracking System', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_heading('API Endpoints - Postman Testing Guide', level=1)
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run(f'\nBase URL: {BASE}\n\nTotal Endpoints: 46')
run.font.size = Pt(14)
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# TABLE OF CONTENTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('Table of Contents', level=1)
toc_items = [
    '1. Setup & Authentication',
    '2. Auth Endpoints (3)',
    '3. User Management Endpoints (11)',
    '4. Division Endpoints (7)',
    '5. Department & OKR Endpoints (14)',
    '6. Evaluation Endpoints (7)',
    '7. Score Level Endpoints (3)',
    '8. Demo Data Endpoint (1)',
    '9. Enum Reference',
    '10. Recommended Testing Order',
]
for item in toc_items:
    doc.add_paragraph(item)
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 1. SETUP & AUTHENTICATION
# ════════════════════════════════════════════════════════════════════
doc.add_heading('1. Setup & Authentication', level=1)
doc.add_paragraph(
    'Before testing any endpoint, you need to obtain a JWT token by logging in. '
    'All endpoints except POST /api/auth/login require a Bearer token in the Authorization header.'
)
doc.add_heading('Postman Setup Steps:', level=2)
steps = [
    'Create a new Postman Environment called "OKR Local".',
    f'Add a variable: baseUrl = {BASE}',
    'Add a variable: token = (leave empty for now).',
    'Send the POST /api/auth/login request (see Section 2).',
    'Copy the token from the response.',
    'Set the token variable value to the JWT token.',
    'For all subsequent requests, go to Authorization tab > Type: Bearer Token > Token: {{token}}.',
]
for i, step in enumerate(steps, 1):
    doc.add_paragraph(f'{i}. {step}')

doc.add_heading('Auto-set Token (Postman Script):', level=2)
doc.add_paragraph(
    'In the Login request\'s "Tests" tab, add the following script to auto-save the token:'
)
add_json_block('Script', 'var jsonData = pm.response.json();\npm.environment.set("token", jsonData.token);')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 2. AUTH ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('2. Auth Endpoints', level=1)

# 2.1 Login
doc.add_heading('2.1  POST /api/auth/login', level=2)
doc.add_paragraph('Description: Authenticate user and receive JWT token.')
doc.add_paragraph('Authorization: None required')
add_table(
    ['Field', 'URL'],
    [['Method', 'POST'], ['URL', f'{BASE}/api/auth/login']]
)
add_json_block('Request Body', '''{
    "username": "admin",
    "password": "admin123"
}''')
add_json_block('Expected Response', '''{
    "token": "eyJhbGciOiJIUzI1NiIs...",
    "user": { ... }
}''')

# 2.2 Register
doc.add_heading('2.2  POST /api/auth/register', level=2)
doc.add_paragraph('Description: Register a new user (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role required)')
add_table(
    ['Field', 'URL'],
    [['Method', 'POST'], ['URL', f'{BASE}/api/auth/register']]
)
add_json_block('Request Body', '''{
    "username": "newuser",
    "email": "newuser@example.com",
    "password": "password123",
    "fullName": "New User",
    "role": "EMPLOYEE",
    "departmentId": "<department-uuid>"
}''')

# 2.3 Get Current User
doc.add_heading('2.3  GET /api/auth/me', level=2)
doc.add_paragraph('Description: Get the currently authenticated user\'s info.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(
    ['Field', 'URL'],
    [['Method', 'GET'], ['URL', f'{BASE}/api/auth/me']]
)
doc.add_paragraph('Request Body: None')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 3. USER MANAGEMENT ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('3. User Management Endpoints', level=1)

# 3.1
doc.add_heading('3.1  GET /api/users', level=2)
doc.add_paragraph('Description: Get all users.')
doc.add_paragraph('Authorization: Bearer Token (ADMIN or DIRECTOR role)')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/users']])
doc.add_paragraph('Request Body: None')

# 3.2
doc.add_heading('3.2  GET /api/users/with-scores', level=2)
doc.add_paragraph('Description: Get all users with their overall scores (Team Overview).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN or DIRECTOR role)')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/users/with-scores']])
doc.add_paragraph('Request Body: None')

# 3.3
doc.add_heading('3.3  GET /api/users/{{id}}', level=2)
doc.add_paragraph('Description: Get user by ID. ADMIN/DIRECTOR can view any user; others can only view themselves.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/users/{{{{id}}}}']])
doc.add_paragraph('Path Parameter: id - UUID of the user')

# 3.4
doc.add_heading('3.4  POST /api/users', level=2)
doc.add_paragraph('Description: Create a new user (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/users']])
add_json_block('Request Body', '''{
    "username": "john_doe",
    "email": "john@example.com",
    "password": "securePass123",
    "fullName": "John Doe",
    "role": "EMPLOYEE",
    "assignedDepartmentIds": ["<dept-uuid-1>", "<dept-uuid-2>"],
    "jobTitle": "Software Engineer",
    "phoneNumber": "+998901234567",
    "bio": "Senior developer"
}''')

# 3.5
doc.add_heading('3.5  PUT /api/users/{{id}}', level=2)
doc.add_paragraph('Description: Update user. ADMIN can update all fields; regular users can only update their own profile fields.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/users/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "fullName": "John Updated",
    "email": "john_updated@example.com",
    "jobTitle": "Lead Engineer",
    "phoneNumber": "+998901234567",
    "bio": "Updated bio",
    "role": "DEPARTMENT_LEADER",
    "assignedDepartmentIds": ["<dept-uuid>"],
    "isActive": true,
    "canEditAssignedDepartments": true,
    "password": "newPassword123"
}''')

# 3.6
doc.add_heading('3.6  DELETE /api/users/{{id}}', level=2)
doc.add_paragraph('Description: Delete a user (Admin only). Cannot delete yourself.')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/users/{{{{id}}}}']])
doc.add_paragraph('Path Parameter: id - UUID of the user to delete')

# 3.7
doc.add_heading('3.7  POST /api/users/{{id}}/departments', level=2)
doc.add_paragraph('Description: Assign departments to a user (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/users/{{{{id}}}}/departments']])
add_json_block('Request Body', '''{
    "departmentIds": ["<dept-uuid-1>", "<dept-uuid-2>"]
}''')

# 3.8
doc.add_heading('3.8  DELETE /api/users/{{id}}/departments/{{deptId}}', level=2)
doc.add_paragraph('Description: Remove a department from a user (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/users/{{{{id}}}}/departments/{{{{deptId}}}}']])
doc.add_paragraph('Path Parameters: id - user UUID, deptId - department UUID')

# 3.9
doc.add_heading('3.9  POST /api/users/{{id}}/photo', level=2)
doc.add_paragraph('Description: Upload a profile photo. ADMIN can upload for any user; users can upload their own.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/users/{{{{id}}}}/photo']])
doc.add_paragraph('Content-Type: multipart/form-data')
add_json_block('Form Data', 'Key: photo\nValue: [Select File]\nType: File')

# 3.10
doc.add_heading('3.10  GET /api/users/by-department/{{deptId}}', level=2)
doc.add_paragraph('Description: Get all users belonging to a specific department.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/users/by-department/{{{{deptId}}}}']])
doc.add_paragraph('Path Parameter: deptId - UUID of the department')

# 3.11
doc.add_heading('3.11  GET /api/users/me/profile', level=2)
doc.add_paragraph('Description: Get the current user\'s extended profile.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/users/me/profile']])
doc.add_paragraph('Request Body: None')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 4. DIVISION ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('4. Division Endpoints', level=1)

# 4.1
doc.add_heading('4.1  GET /api/divisions', level=2)
doc.add_paragraph('Description: Get all divisions.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/divisions']])
doc.add_paragraph('Request Body: None')

# 4.2
doc.add_heading('4.2  GET /api/divisions/{{id}}', level=2)
doc.add_paragraph('Description: Get a single division by ID.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/divisions/{{{{id}}}}']])
doc.add_paragraph('Path Parameter: id - UUID of the division')

# 4.3
doc.add_heading('4.3  POST /api/divisions', level=2)
doc.add_paragraph('Description: Create a new division (Admin/Director only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN or DIRECTOR role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/divisions']])
add_json_block('Request Body', '''{
    "name": "Technology Division",
    "leaderId": "<user-uuid>"
}''')

# 4.4
doc.add_heading('4.4  PUT /api/divisions/{{id}}', level=2)
doc.add_paragraph('Description: Update a division (requires edit permission).')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/divisions/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "name": "Updated Division Name",
    "leaderId": "<user-uuid>"
}''')

# 4.5
doc.add_heading('4.5  DELETE /api/divisions/{{id}}', level=2)
doc.add_paragraph('Description: Delete a division (Admin only). Fails if division has departments.')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/divisions/{{{{id}}}}']])
doc.add_paragraph('Path Parameter: id - UUID of the division')

# 4.6
doc.add_heading('4.6  GET /api/divisions/{{id}}/departments', level=2)
doc.add_paragraph('Description: Get all departments within a division.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/divisions/{{{{id}}}}/departments']])
doc.add_paragraph('Path Parameter: id - UUID of the division')

# 4.7
doc.add_heading('4.7  GET /api/divisions/{{id}}/score', level=2)
doc.add_paragraph('Description: Get division with calculated scores.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/divisions/{{{{id}}}}/score']])
doc.add_paragraph('Path Parameter: id - UUID of the division')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 5. DEPARTMENT & OKR ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('5. Department & OKR Endpoints', level=1)

# Departments
doc.add_heading('Departments', level=2)

# 5.1
doc.add_heading('5.1  GET /api/departments', level=3)
doc.add_paragraph('Description: Get all departments.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/departments']])
doc.add_paragraph('Request Body: None')

# 5.2
doc.add_heading('5.2  GET /api/departments/{{id}}', level=3)
doc.add_paragraph('Description: Get a single department by ID with its objectives and key results.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/departments/{{{{id}}}}']])

# 5.3
doc.add_heading('5.3  POST /api/departments', level=3)
doc.add_paragraph('Description: Create a new department (Admin/Director only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN or DIRECTOR role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/departments']])
add_json_block('Request Body', '''{
    "name": "Engineering Department",
    "divisionId": "<division-uuid>"
}''')

# 5.4
doc.add_heading('5.4  PUT /api/departments/{{id}}', level=3)
doc.add_paragraph('Description: Update a department (requires edit permission).')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/departments/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "name": "Updated Department Name",
    "divisionId": "<division-uuid>"
}''')

# 5.5
doc.add_heading('5.5  DELETE /api/departments/{{id}}', level=3)
doc.add_paragraph('Description: Delete a department (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/departments/{{{{id}}}}']])

# 5.6
doc.add_heading('5.6  GET /api/departments/{{id}}/scores', level=3)
doc.add_paragraph('Description: Get department scores with evaluations.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/departments/{{{{id}}}}/scores']])

# Objectives
doc.add_heading('Objectives', level=2)

# 5.7
doc.add_heading('5.7  POST /api/departments/{{departmentId}}/objectives', level=3)
doc.add_paragraph('Description: Create an objective within a department.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/departments/{{{{departmentId}}}}/objectives']])
add_json_block('Request Body', '''{
    "name": "Increase Revenue",
    "weight": 40
}''')

# 5.8
doc.add_heading('5.8  PUT /api/objectives/{{id}}', level=3)
doc.add_paragraph('Description: Update an objective.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/objectives/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "name": "Increase Revenue by 20%",
    "weight": 50
}''')

# 5.9
doc.add_heading('5.9  DELETE /api/objectives/{{id}}', level=3)
doc.add_paragraph('Description: Delete an objective.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/objectives/{{{{id}}}}']])

# Key Results
doc.add_heading('Key Results', level=2)

# 5.10
doc.add_heading('5.10  POST /api/objectives/{{objectiveId}}/key-results', level=3)
doc.add_paragraph('Description: Create a key result within an objective.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/objectives/{{{{objectiveId}}}}/key-results']])
add_json_block('Request Body', '''{
    "name": "Monthly recurring revenue",
    "description": "Track MRR growth",
    "metricType": "HIGHER_BETTER",
    "unit": "USD",
    "weight": 50,
    "thresholds": {
        "min": "10000",
        "target": "50000",
        "stretch": "75000"
    },
    "actualValue": "0"
}''')

# 5.11
doc.add_heading('5.11  PUT /api/key-results/{{id}}', level=3)
doc.add_paragraph('Description: Update a key result.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/key-results/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "name": "Updated key result name",
    "description": "Updated description",
    "metricType": "HIGHER_BETTER",
    "unit": "USD",
    "weight": 60,
    "thresholds": {
        "min": "15000",
        "target": "60000",
        "stretch": "80000"
    },
    "actualValue": "25000"
}''')

# 5.12
doc.add_heading('5.12  PUT /api/key-results/{{id}}/actual-value', level=3)
doc.add_paragraph('Description: Update only the actual value of a key result.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/key-results/{{{{id}}}}/actual-value']])
add_json_block('Request Body', '''{
    "actualValue": "35000"
}''')

# 5.13
doc.add_heading('5.13  DELETE /api/key-results/{{id}}', level=3)
doc.add_paragraph('Description: Delete a key result.')
doc.add_paragraph('Authorization: Bearer Token (requires department edit permission)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/key-results/{{{{id}}}}']])

# Export
doc.add_heading('Export', level=2)

# 5.14
doc.add_heading('5.14  GET /api/export/excel', level=3)
doc.add_paragraph('Description: Export all OKR data to an Excel file (.xlsx).')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/export/excel']])
doc.add_paragraph('Request Body: None')
doc.add_paragraph('Response: File download (application/vnd.openxmlformats-officedocument.spreadsheetml.sheet)')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 6. EVALUATION ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('6. Evaluation Endpoints', level=1)

# 6.1
doc.add_heading('6.1  POST /api/evaluations', level=2)
doc.add_paragraph('Description: Create a new evaluation (draft).')
doc.add_paragraph('Authorization: Bearer Token (DIRECTOR, HR, BUSINESS_BLOCK, or ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/evaluations']])
add_json_block('Request Body (DIRECTOR evaluation)', '''{
    "targetType": "DEPARTMENT",
    "targetId": "<department-uuid>",
    "evaluatorType": "DIRECTOR",
    "starRating": 4,
    "comment": "Good performance this quarter"
}''')
add_json_block('Request Body (HR evaluation)', '''{
    "targetType": "DEPARTMENT",
    "targetId": "<department-uuid>",
    "evaluatorType": "HR",
    "letterRating": "A",
    "comment": "Excellent team collaboration"
}''')
add_json_block('Request Body (BUSINESS_BLOCK evaluation)', '''{
    "targetType": "DEPARTMENT",
    "targetId": "<department-uuid>",
    "evaluatorType": "BUSINESS_BLOCK",
    "numericRating": 4.5,
    "comment": "Strong business impact"
}''')

# 6.2
doc.add_heading('6.2  POST /api/evaluations/{{id}}/submit', level=2)
doc.add_paragraph('Description: Submit a draft evaluation (changes status from DRAFT to SUBMITTED).')
doc.add_paragraph('Authorization: Bearer Token (DIRECTOR, HR, BUSINESS_BLOCK, or ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/evaluations/{{{{id}}}}/submit']])
doc.add_paragraph('Request Body: None')

# 6.3
doc.add_heading('6.3  PUT /api/evaluations/{{id}}', level=2)
doc.add_paragraph('Description: Update an existing evaluation.')
doc.add_paragraph('Authorization: Bearer Token (DIRECTOR, HR, BUSINESS_BLOCK, or ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/evaluations/{{{{id}}}}']])
add_json_block('Request Body', '''{
    "targetType": "DEPARTMENT",
    "targetId": "<department-uuid>",
    "evaluatorType": "DIRECTOR",
    "starRating": 5,
    "comment": "Updated: Outstanding performance"
}''')

# 6.4
doc.add_heading('6.4  GET /api/evaluations/target/{{type}}/{{id}}', level=2)
doc.add_paragraph('Description: Get all evaluations for a specific target (department or employee).')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/evaluations/target/{{{{type}}}}/{{{{id}}}}']])
doc.add_paragraph('Path Parameters:')
doc.add_paragraph('  type - "DEPARTMENT" or "EMPLOYEE"')
doc.add_paragraph('  id - UUID of the target')

# 6.5
doc.add_heading('6.5  GET /api/evaluations/my', level=2)
doc.add_paragraph('Description: Get all evaluations created by the current user.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/evaluations/my']])
doc.add_paragraph('Request Body: None')

# 6.6
doc.add_heading('6.6  GET /api/evaluations/all', level=2)
doc.add_paragraph('Description: Get all evaluations in the system (Admin debug endpoint).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/evaluations/all']])
doc.add_paragraph('Request Body: None')

# 6.7
doc.add_heading('6.7  DELETE /api/evaluations/{{id}}', level=2)
doc.add_paragraph('Description: Delete a draft evaluation.')
doc.add_paragraph('Authorization: Bearer Token (DIRECTOR, HR, BUSINESS_BLOCK, or ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'DELETE'], ['URL', f'{BASE}/api/evaluations/{{{{id}}}}']])
doc.add_paragraph('Path Parameter: id - UUID of the evaluation')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 7. SCORE LEVEL ENDPOINTS
# ════════════════════════════════════════════════════════════════════
doc.add_heading('7. Score Level Endpoints', level=1)

# 7.1
doc.add_heading('7.1  GET /api/score-levels', level=2)
doc.add_paragraph('Description: Get all score levels.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'GET'], ['URL', f'{BASE}/api/score-levels']])
doc.add_paragraph('Request Body: None')

# 7.2
doc.add_heading('7.2  PUT /api/score-levels', level=2)
doc.add_paragraph('Description: Update score levels (replaces all).')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'PUT'], ['URL', f'{BASE}/api/score-levels']])
add_json_block('Request Body', '''[
    {
        "id": "<existing-id-or-null>",
        "name": "Excellent",
        "scoreValue": 5.0,
        "color": "#22C55E",
        "displayOrder": 1
    },
    {
        "id": "<existing-id-or-null>",
        "name": "Good",
        "scoreValue": 4.0,
        "color": "#3B82F6",
        "displayOrder": 2
    },
    {
        "id": "<existing-id-or-null>",
        "name": "Average",
        "scoreValue": 3.0,
        "color": "#F59E0B",
        "displayOrder": 3
    }
]''')

# 7.3
doc.add_heading('7.3  POST /api/score-levels/reset', level=2)
doc.add_paragraph('Description: Reset score levels to default values.')
doc.add_paragraph('Authorization: Bearer Token')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/score-levels/reset']])
doc.add_paragraph('Request Body: None')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 8. DEMO DATA ENDPOINT
# ════════════════════════════════════════════════════════════════════
doc.add_heading('8. Demo Data Endpoint', level=1)

doc.add_heading('8.1  POST /api/demo/load', level=2)
doc.add_paragraph('Description: Load demo/sample data into the system (Admin only).')
doc.add_paragraph('Authorization: Bearer Token (ADMIN role)')
add_table(['Field', 'Value'], [['Method', 'POST'], ['URL', f'{BASE}/api/demo/load']])
doc.add_paragraph('Request Body: None')
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 9. ENUM REFERENCE
# ════════════════════════════════════════════════════════════════════
doc.add_heading('9. Enum Reference', level=1)

doc.add_heading('Role', level=2)
add_table(
    ['Value', 'Description'],
    [
        ['EMPLOYEE', 'Regular employee - can view own OKRs'],
        ['DEPARTMENT_LEADER', 'Department leader - manages department OKRs'],
        ['HR', 'HR staff - evaluates with letter grades (A-D)'],
        ['DIRECTOR', 'Top leadership - evaluates with star ratings (1-5)'],
        ['BUSINESS_BLOCK', 'Business block leaders - evaluates with 1-5 numeric rating'],
        ['ADMIN', 'System administrator - full access'],
    ]
)

doc.add_heading('EvaluatorType', level=2)
add_table(
    ['Value', 'Rating Scale', 'Weight'],
    [
        ['DIRECTOR', '1-5 stars (mapped to 4.25-5.0)', '20% of final score'],
        ['HR', 'A, B, C, D letters (A=5.0, B=4.75, C=4.5, D=4.25)', '20% of final score'],
        ['BUSINESS_BLOCK', '1-5 numeric rating', 'Separate (not in weighted score)'],
    ]
)

doc.add_heading('MetricType (Key Results)', level=2)
add_table(
    ['Value', 'Description'],
    [
        ['HIGHER_BETTER', 'Higher actual value = better score (e.g., revenue)'],
        ['LOWER_BETTER', 'Lower actual value = better score (e.g., bug count)'],
        ['QUALITATIVE', 'Non-numeric / qualitative assessment'],
    ]
)
doc.add_page_break()

# ════════════════════════════════════════════════════════════════════
# 10. RECOMMENDED TESTING ORDER
# ════════════════════════════════════════════════════════════════════
doc.add_heading('10. Recommended Testing Order', level=1)

doc.add_paragraph('Follow this order to test endpoints logically, creating dependencies before they are needed:')

steps = [
    ('Step 1: Authentication', [
        'POST /api/auth/login - Get JWT token',
        'GET /api/auth/me - Verify token works',
    ]),
    ('Step 2: Load Demo Data (optional)', [
        'POST /api/demo/load - Populates sample divisions, departments, objectives, key results',
    ]),
    ('Step 3: Divisions (CRUD)', [
        'GET /api/divisions - List all',
        'POST /api/divisions - Create one',
        'GET /api/divisions/{id} - Get created division',
        'PUT /api/divisions/{id} - Update it',
        'GET /api/divisions/{id}/score - Check scores',
    ]),
    ('Step 4: Departments (CRUD)', [
        'GET /api/departments - List all',
        'POST /api/departments - Create one (use divisionId from Step 3)',
        'GET /api/departments/{id} - Get created department',
        'PUT /api/departments/{id} - Update it',
        'GET /api/departments/{id}/scores - Check scores',
        'GET /api/divisions/{id}/departments - Verify department appears in division',
    ]),
    ('Step 5: Users', [
        'GET /api/users - List all users',
        'POST /api/users - Create a test user',
        'GET /api/users/{id} - Get created user',
        'PUT /api/users/{id} - Update user',
        'POST /api/users/{id}/departments - Assign departments',
        'GET /api/users/by-department/{deptId} - Verify user appears',
        'GET /api/users/with-scores - Check team overview',
        'GET /api/users/me/profile - Check current user profile',
    ]),
    ('Step 6: Objectives', [
        'POST /api/departments/{deptId}/objectives - Create objective',
        'PUT /api/objectives/{id} - Update objective',
    ]),
    ('Step 7: Key Results', [
        'POST /api/objectives/{objId}/key-results - Create key result',
        'PUT /api/key-results/{id} - Update key result',
        'PUT /api/key-results/{id}/actual-value - Update actual value',
    ]),
    ('Step 8: Evaluations', [
        'POST /api/evaluations - Create evaluation (try all 3 evaluator types)',
        'GET /api/evaluations/my - Check your evaluations',
        'POST /api/evaluations/{id}/submit - Submit evaluation',
        'GET /api/evaluations/target/{type}/{id} - View evaluations for target',
        'GET /api/evaluations/all - Admin: view all evaluations',
    ]),
    ('Step 9: Score Levels', [
        'GET /api/score-levels - View current levels',
        'PUT /api/score-levels - Update levels',
        'POST /api/score-levels/reset - Reset to defaults',
    ]),
    ('Step 10: Export', [
        'GET /api/export/excel - Download Excel report',
    ]),
    ('Step 11: Cleanup (DELETE operations)', [
        'DELETE /api/key-results/{id}',
        'DELETE /api/objectives/{id}',
        'DELETE /api/evaluations/{id}',
        'DELETE /api/users/{id}/departments/{deptId}',
        'DELETE /api/users/{id}',
        'DELETE /api/departments/{id}',
        'DELETE /api/divisions/{id}',
    ]),
]

for title, items in steps:
    doc.add_heading(title, level=2)
    for item in items:
        doc.add_paragraph(item, style='List Bullet')

# ── Save ────────────────────────────────────────────────────────────
output_path = '/Users/spencercoder/IdeaProjects/OKR/OKR_API_Postman_Guide.docx'
doc.save(output_path)
print(f'Document saved to: {output_path}')
